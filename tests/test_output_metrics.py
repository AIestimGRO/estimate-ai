"""Regression tests for price metrics in generated WA files."""

from types import SimpleNamespace

from openpyxl import Workbook

from core.excel_writer import (
    GlobalAnalogPlan,
    WriterColumns,
    _plan_output_columns,
    _write_metric_formulas,
    _write_metric_headers,
)


def _columns():
    layout = WriterColumns(
        base_price=6, average=7, code=14, code_kr=15,
        section=16, analog_start=20, header_row=7,
    )
    return _plan_output_columns(layout, SimpleNamespace(column=7), 20)


def _plan(last):
    return GlobalAnalogPlan(
        by_key={}, tkp_winner_by_key={}, tkp_reserve_by_key={},
        columns=(), last_column=last,
    )


def test_metrics_after_section_and_before_analogs():
    cols = _columns()
    assert (cols.deviation, cols.min_price, cols.max_price, cols.analog_start) == (17,18,19,20)
    sheet = Workbook().active
    _write_metric_headers(sheet, 7, cols)
    assert "Отклонение" in sheet["Q7"].value
    assert "Мин." in sheet["R7"].value
    assert "Макс." in sheet["S7"].value


def test_dynamic_formulas_cover_rnmc_and_paired_tkp_columns():
    sheet = Workbook().active
    _write_metric_formulas(sheet, 9, _columns(), _plan(23))
    assert sheet["Q9"].value == '=IFERROR(IF(OR(NOT(ISNUMBER(F9)),F9=0),"",(G9-F9)/F9),"")'
    assert sheet["Q9"].number_format == "+0.00%;-0.00%;0.00%"
    assert sheet["R9"].value == '=IF(COUNTIF(T9:W9,">0")=0,"",MIN(T9:W9))'
    assert sheet["S9"].value == '=IF(COUNTIF(T9:W9,">0")=0,"",MAX(T9:W9))'


def test_price_metrics_blank_without_candidates():
    sheet = Workbook().active
    _write_metric_formulas(sheet, 9, _columns(), _plan(19))
    assert sheet["R9"].value is None
    assert sheet["S9"].value is None



def test_minimum_formula_in_xlsx_is_standard_min(tmp_path):
    """The saved XLSX must contain ordinary MIN/MAX, not future functions."""
    from zipfile import ZipFile

    workbook = Workbook()
    sheet = workbook.active
    _write_metric_formulas(sheet, 9, _columns(), _plan(23))
    path = tmp_path / "metrics.xlsx"
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "MIN(T9:W9)" in xml
    assert "MAX(T9:W9)" in xml
    assert "_xlfn." not in xml
    assert "AGGREGATE" not in xml


def test_minimum_and_maximum_formulas_survive_deleted_candidates():
    """Both summary cells remain formulas after individual analogs are removed."""
    sheet = Workbook().active
    for col, price in (("T", 100), ("U", 150), ("V", 180)):
        sheet[f"{col}9"] = price
    _write_metric_formulas(sheet, 9, _columns(), _plan(22))
    expected_min = sheet["R9"].value
    expected_max = sheet["S9"].value
    sheet["T9"] = None
    sheet["V9"] = None
    assert sheet["R9"].value == expected_min
    assert sheet["S9"].value == expected_max
    assert sheet["U9"].value == 150
